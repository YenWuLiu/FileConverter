import csv
import shutil
import subprocess
import tempfile
from pathlib import Path

from .. import deps
from .base import ConvertError, register
from .textdata import _read_text_auto


@register("docx", ["txt"], requires=("docx",), label="Word→纯文本")
def docx_to_txt(src: Path, dst: Path, **opts):
    import docx

    document = docx.Document(str(src))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    dst.write_text("\n".join(lines), encoding="utf-8")


@register("docx", ["md"], requires=("docx",), label="Word→Markdown")
def docx_to_md(src: Path, dst: Path, **opts):
    import docx

    document = docx.Document(str(src))
    lines = []
    for p in document.paragraphs:
        style = (p.style.name or "") if p.style else ""
        text = p.text
        if style.startswith("Heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            lines.append("#" * min(max(level, 1), 6) + " " + text)
        else:
            lines.append(text)
    dst.write_text("\n\n".join(lines) + "\n", encoding="utf-8")


@register("xlsx", ["csv"], requires=("openpyxl",), label="Excel→CSV")
def xlsx_to_csv(src: Path, dst: Path, **opts):
    import openpyxl

    from ..registry import unique_path

    wb = openpyxl.load_workbook(str(src), data_only=True)
    names = wb.sheetnames
    for name in names:
        ws = wb[name]
        target = dst if len(names) == 1 else unique_path(dst.with_name(f"{dst.stem}_{name}{dst.suffix}"))
        with open(target, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])


@register("csv", ["xlsx"], requires=("openpyxl",), label="CSV→Excel")
def csv_to_xlsx(src: Path, dst: Path, src_encoding=None, **opts):
    import openpyxl

    text = _read_text_auto(src, src_encoding)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = src.stem[:31] or "Sheet1"
    for row in csv.reader(text.splitlines()):
        ws.append(row)
    wb.save(str(dst))


@register("pptx", ["txt"], requires=("pptx",), label="PPT→纯文本")
def pptx_to_txt(src: Path, dst: Path, **opts):
    from pptx import Presentation

    prs = Presentation(str(src))
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"=== 第 {i} 页 ===")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs)
                    if t.strip():
                        lines.append(t)
    dst.write_text("\n".join(lines) + "\n", encoding="utf-8")


_COM_APPS = {
    ".docx": "Word.Application",
    ".xlsx": "Excel.Application",
    ".pptx": "PowerPoint.Application",
}


def _com_to_pdf(src: Path, dst: Path) -> None:
    import pythoncom
    import win32com.client

    ext = src.suffix.lower()
    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx(_COM_APPS[ext])
        if ext != ".pptx":  # PowerPoint 不允许 Visible=False
            app.Visible = False
        if ext == ".docx":
            doc = app.Documents.Open(str(src.resolve()), ReadOnly=True)
            doc.SaveAs2(str(dst.resolve()), FileFormat=17)  # wdFormatPDF
            doc.Close(False)
        elif ext == ".xlsx":
            wb = app.Workbooks.Open(str(src.resolve()), ReadOnly=True)
            wb.ExportAsFixedFormat(0, str(dst.resolve()))  # xlTypePDF = 0
            wb.Close(False)
        else:
            pres = app.Presentations.Open(str(src.resolve()), WithWindow=False)
            pres.SaveAs(str(dst.resolve()), 32)  # ppSaveAsPDF
            pres.Close()
    except Exception as e:
        raise ConvertError(f"Office 转换 PDF 失败 {src.name}: {e}")
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def _libreoffice_to_pdf(src: Path, dst: Path) -> None:
    exe = deps.libreoffice_path()
    dst.parent.mkdir(parents=True, exist_ok=True)
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            r = subprocess.run(
                [exe, "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(src)],
                capture_output=True, text=True, timeout=300,
            )
            produced = Path(tmpdir) / (src.stem + ".pdf")
            if r.returncode != 0 or not produced.exists():
                raise ConvertError(f"LibreOffice 转换失败 {src.name}: {(r.stderr or r.stdout).strip()}")
            shutil.move(str(produced), str(dst))
    except subprocess.TimeoutExpired:
        raise ConvertError(f"LibreOffice 转换超时 {src.name}")


@register(["docx", "xlsx", "pptx"], ["pdf"], label="Office→PDF")
def office_to_pdf(src: Path, dst: Path, **opts):
    if deps.office_available():
        try:
            _com_to_pdf(src, dst)
        except ConvertError:
            if deps.libreoffice_path():
                _libreoffice_to_pdf(src, dst)
            else:
                raise
    elif deps.libreoffice_path():
        _libreoffice_to_pdf(src, dst)
    else:
        raise ConvertError("转换为 PDF 需要安装 MS Office 或 LibreOffice")


def _pdf_to_docx_com(src: Path, dst: Path) -> None:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx("Word.Application")
        app.Visible = False
        app.DisplayAlerts = 0  # wdAlertsNone
        doc = app.Documents.Open(str(src.resolve()), ReadOnly=True)
        doc.SaveAs2(str(dst.resolve()), FileFormat=16)  # wdFormatDocumentDefault (.docx)
        doc.Close(False)
    except Exception as e:
        raise ConvertError(f"Word 转换 PDF 失败 {src.name}: {e}")
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


@register("pdf", ["docx"], label="PDF→Word")
def pdf_to_docx(src: Path, dst: Path, **opts):
    lib_err: Exception | None = None
    if deps.module_available("pdf2docx"):
        from pdf2docx import Converter

        cv = Converter(str(src))
        try:
            cv.convert(str(dst))
        except Exception as e:  # 只有 convert 失败才清理并回退 COM
            lib_err = e
            if dst.exists():  # 清掉半成品，避免 Word 覆盖提示
                dst.unlink()
        finally:
            cv.close()  # close 异常直接外抛，不会误删已生成的有效产物
        if lib_err is None:
            return
    if deps.office_available():
        _pdf_to_docx_com(src, dst)
        return
    if lib_err is not None:
        raise ConvertError(f"PDF 转 Word 失败 {src.name}: {lib_err}") from lib_err
    raise ConvertError("PDF 转 Word 需要 pdf2docx 库或 MS Office")
