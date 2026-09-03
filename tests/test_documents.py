import csv
import subprocess
from pathlib import Path

import pytest

from file_converter import registry


def test_docx_to_txt_and_md(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("报告标题", level=1)
    d.add_paragraph("第一段内容")
    p = tmp_path / "a.docx"
    d.save(str(p))

    txt = registry.convert(p, "txt")
    assert "第一段内容" in txt.read_text(encoding="utf-8")

    md = registry.convert(p, "md")
    md_text = md.read_text(encoding="utf-8")
    assert md_text.startswith("# 报告标题")


def test_xlsx_to_csv(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "分数"])
    ws.append(["张三", 90])
    p = tmp_path / "a.xlsx"
    wb.save(str(p))

    out = registry.convert(p, "csv")
    rows = list(csv.reader(out.read_text(encoding="utf-8-sig").splitlines()))
    assert rows == [["姓名", "分数"], ["张三", "90"]]


def test_xlsx_to_csv_multisheet(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "一"
    wb.create_sheet("二")
    wb["一"].append(["a"])
    wb["二"].append(["b"])
    p = tmp_path / "m.xlsx"
    wb.save(str(p))

    registry.convert(p, "csv")
    assert (tmp_path / "m_一.csv").exists()
    assert (tmp_path / "m_二.csv").exists()


def test_xlsx_to_csv_multisheet_no_overwrite(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "一"
    wb.create_sheet("二")
    p = tmp_path / "m.xlsx"
    wb.save(str(p))
    registry.convert(p, "csv")
    registry.convert(p, "csv")
    assert (tmp_path / "m_一.csv").exists()
    assert (tmp_path / "m_一 (1).csv").exists()


def test_csv_to_xlsx(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("x,y\n1,2\n", encoding="utf-8-sig")
    out = registry.convert(src, "xlsx")

    import openpyxl

    wb = openpyxl.load_workbook(str(out))
    ws = wb.active
    assert ws["A1"].value == "x" and ws["B2"].value == "2"


def test_pptx_to_txt(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "演示标题"
    slide.placeholders[1].text = "要点一"
    p = tmp_path / "a.pptx"
    prs.save(str(p))

    out = registry.convert(p, "txt")
    text = out.read_text(encoding="utf-8")
    assert "演示标题" in text and "要点一" in text


def test_pdf_to_docx(tmp_path):
    import docx
    import fitz

    p = tmp_path / "a.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_font(fontname="cjk", fontfile="C:/Windows/Fonts/simsun.ttc")
    page.insert_text((72, 72), "PDF转Word测试文档", fontname="cjk")
    pdf.save(str(p))
    pdf.close()

    out = registry.convert(p, "docx")
    d = docx.Document(str(out))
    text = "\n".join(par.text for par in d.paragraphs)
    assert "PDF转Word测试文档" in text


class _FailingConverter:
    """假 pdf2docx.Converter：convert 抛异常。"""

    def __init__(self, path):
        pass

    def convert(self, out):
        raise RuntimeError("引擎内部错误")

    def close(self):
        pass


def test_pdf_to_docx_pdf2docx_failure_no_office_reports_real_error(tmp_path, monkeypatch):
    import pdf2docx

    from file_converter import deps
    from file_converter.converters.base import ConvertError

    p = tmp_path / "a.pdf"
    p.write_bytes(b"%PDF-fake")

    monkeypatch.setattr(pdf2docx, "Converter", _FailingConverter)
    monkeypatch.setattr(deps, "office_available", lambda: False)

    with pytest.raises(ConvertError, match="PDF 转 Word 失败 a.pdf"):
        registry.convert(p, "docx")


def test_pdf_to_docx_pdf2docx_failure_falls_back_to_com(tmp_path, monkeypatch):
    import pdf2docx

    from file_converter import deps
    from file_converter.converters import documents

    p = tmp_path / "a.pdf"
    p.write_bytes(b"%PDF-fake")

    monkeypatch.setattr(pdf2docx, "Converter", _FailingConverter)
    monkeypatch.setattr(deps, "office_available", lambda: True)
    monkeypatch.setattr(documents, "_pdf_to_docx_com", lambda s, d: d.write_bytes(b"COM-OUT"))

    out = registry.convert(p, "docx")
    assert out.read_bytes() == b"COM-OUT"


def test_pdf_to_docx_close_failure_keeps_valid_output(tmp_path, monkeypatch):
    import pdf2docx

    from file_converter import deps
    from file_converter.converters.base import ConvertError

    class CloseBoomConverter:
        def __init__(self, path):
            pass

        def convert(self, out):
            Path(out).write_bytes(b"OK")

        def close(self):
            raise RuntimeError("close boom")

    p = tmp_path / "a.pdf"
    p.write_bytes(b"%PDF-fake")

    monkeypatch.setattr(pdf2docx, "Converter", CloseBoomConverter)
    monkeypatch.setattr(deps, "office_available", lambda: False)

    with pytest.raises(ConvertError):
        registry.convert(p, "docx")
    assert (tmp_path / "a.docx").read_bytes() == b"OK"


def test_libreoffice_to_pdf_does_not_overwrite_existing(tmp_path, monkeypatch):
    import docx

    from file_converter import deps
    from file_converter.converters import documents

    d = docx.Document()
    d.add_paragraph("内容")
    src = tmp_path / "a.docx"
    d.save(str(src))

    sentinel = tmp_path / "a.pdf"
    sentinel.write_bytes(b"SENTINEL-ORIGINAL")

    monkeypatch.setattr(deps, "office_available", lambda: False)
    monkeypatch.setattr(deps, "libreoffice_path", lambda: "soffice")

    def fake_run(cmd, **kwargs):
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        assert outdir != tmp_path
        (outdir / "a.pdf").write_bytes(b"FAKE-PDF")
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(documents.subprocess, "run", fake_run)

    dst = registry.unique_path(tmp_path / "a.pdf")
    assert dst != sentinel
    out = registry.convert(src, "pdf", dst)
    assert out == dst
    assert sentinel.read_bytes() == b"SENTINEL-ORIGINAL"
    assert out.read_bytes() == b"FAKE-PDF"
